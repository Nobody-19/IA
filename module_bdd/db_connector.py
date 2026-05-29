import csv
import os
from sqlalchemy import create_engine, text, inspect
from pathlib import Path

DB_PATH = Path(__file__).parent.parent / "entreprise.db"
engine  = create_engine(f"sqlite:///{DB_PATH}", echo=False)

def executer_sql(sql: str) -> tuple:
    with engine.connect() as conn:
        result  = conn.execute(text(sql))
        columns = list(result.keys())
        rows    = result.fetchall()
        return columns, [list(r) for r in rows]

def get_schema() -> str:
    inspecteur = inspect(engine)
    schema     = []
    for table in inspecteur.get_table_names():
        cols = inspecteur.get_columns(table)
        fks  = inspecteur.get_foreign_keys(table)
        col_defs = ", ".join(f"{c['name']} {c['type']}" for c in cols)
        schema.append(f"{table}({col_defs})")
        for fk in fks:
            schema.append(f"  FK: {fk['constrained_columns']} -> {fk['referred_table']}({fk['referred_columns']})")
    return "\n".join(schema)

def formater_resultats(columns: list, rows: list) -> str:
    if not rows:
        return "Aucun resultat."
    col_width = [
        max(len(str(c)), max((len(str(r[i])) for r in rows), default=0))
        for i, c in enumerate(columns)
    ]
    header    = " | ".join(str(c).ljust(col_width[i]) for i, c in enumerate(columns))
    separator = "-+-".join("-" * w for w in col_width)
    lignes    = [header, separator]
    for row in rows:
        lignes.append(" | ".join(str(v).ljust(col_width[i]) for i, v in enumerate(row)))
    return "\n".join(lignes)

def get_stats() -> dict:
    stats = {}
    with engine.connect() as conn:
        for table in inspect(engine).get_table_names():
            count = conn.execute(text(f"SELECT COUNT(*) FROM {table}")).scalar()
            stats[table] = count
    return stats

def lister_tables() -> list:
    return inspect(engine).get_table_names()

def importer_csv(chemin_csv: str, nom_table: str, separateur: str = ",") -> dict:
    chemin = Path(chemin_csv)
    if not chemin.exists():
        return {"succes": False, "erreur": f"Fichier introuvable : {chemin_csv}"}
    try:
        with open(chemin, encoding="utf-8-sig", errors="ignore") as f:
            reader  = csv.DictReader(f, delimiter=separateur)
            colonnes = reader.fieldnames
            if not colonnes:
                return {"succes": False, "erreur": "Fichier CSV vide ou sans en-tetes"}
            lignes = list(reader)
        if not lignes:
            return {"succes": False, "erreur": "Aucune donnee dans le fichier"}
        col_defs = ", ".join(f'"{c}" TEXT' for c in colonnes)
        with engine.connect() as conn:
            conn.execute(text(f'CREATE TABLE IF NOT EXISTS "{nom_table}" ({col_defs})'))
            conn.execute(text(f'DELETE FROM "{nom_table}"'))
            for ligne in lignes:
                placeholders = ", ".join(f":{i}" for i in range(len(colonnes)))
                vals = {str(i): ligne[c] for i, c in enumerate(colonnes)}
                conn.execute(
                    text(f'INSERT INTO "{nom_table}" VALUES ({placeholders})'),
                    vals
                )
            conn.commit()
        return {
            "succes":   True,
            "table":    nom_table,
            "colonnes": list(colonnes),
            "lignes":   len(lignes)
        }
    except Exception as e:
        return {"succes": False, "erreur": str(e)}

def importer_excel(chemin_excel: str, nom_table: str, feuille: int = 0) -> dict:
    try:
        import openpyxl
    except ImportError:
        return {"succes": False, "erreur": "Installe openpyxl : pip install openpyxl"}
    chemin = Path(chemin_excel)
    if not chemin.exists():
        return {"succes": False, "erreur": f"Fichier introuvable : {chemin_excel}"}
    try:
        wb      = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
        ws      = wb.worksheets[feuille]
        lignes  = list(ws.iter_rows(values_only=True))
        if not lignes:
            return {"succes": False, "erreur": "Feuille vide"}
        colonnes = [str(c) if c else f"col_{i}" for i, c in enumerate(lignes[0])]
        donnees  = lignes[1:]
        col_defs = ", ".join(f'"{c}" TEXT' for c in colonnes)
        with engine.connect() as conn:
            conn.execute(text(f'CREATE TABLE IF NOT EXISTS "{nom_table}" ({col_defs})'))
            conn.execute(text(f'DELETE FROM "{nom_table}"'))
            for row in donnees:
                if all(v is None for v in row):
                    continue
                placeholders = ", ".join(f":{i}" for i in range(len(colonnes)))
                vals = {str(i): str(v) if v is not None else "" for i, v in enumerate(row)}
                conn.execute(
                    text(f'INSERT INTO "{nom_table}" VALUES ({placeholders})'),
                    vals
                )
            conn.commit()
        return {
            "succes":   True,
            "table":    nom_table,
            "colonnes": colonnes,
            "lignes":   len(donnees)
        }
    except Exception as e:
        return {"succes": False, "erreur": str(e)}